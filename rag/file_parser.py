import os
from pathlib import Path
from typing import List, Dict, Any


class FileParser:
    @staticmethod
    def parse(file_path: str, file_type: str) -> List[Dict[str, Any]]:
        """
        Parse file and return list of pages with content and page numbers.
        Returns: [{'content': str, 'page_number': int, 'file_name': str}, ...]
        """
        file_name = Path(file_path).name
        
        if file_type == "pdf":
            return FileParser._parse_pdf(file_path, file_name)
        elif file_type == "docx":
            return FileParser._parse_docx(file_path, file_name)
        elif file_type == "pptx":
            return FileParser._parse_pptx(file_path, file_name)
        elif file_type == "txt":
            return FileParser._parse_txt(file_path, file_name)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    @staticmethod
    def _parse_pdf(file_path: str, file_name: str) -> List[Dict[str, Any]]:
        import fitz
        pages = []
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc, 1):
            text = page.get_text()
            if text.strip():
                pages.append({
                    "content": text,
                    "page_number": page_num,
                    "file_name": file_name,
                })
        doc.close()
        return pages

    @staticmethod
    def _parse_docx(file_path: str, file_name: str) -> List[Dict[str, Any]]:
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        content = "\n\n".join(full_text)
        return [{
            "content": content,
            "page_number": 1,
            "file_name": file_name,
        }]

    @staticmethod
    def _parse_pptx(file_path: str, file_name: str) -> List[Dict[str, Any]]:
        from pptx import Presentation
        prs = Presentation(file_path)
        pages = []
        for slide_num, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        if para.text.strip():
                            texts.append(para.text)
            if texts:
                pages.append({
                    "content": "\n".join(texts),
                    "page_number": slide_num,
                    "file_name": file_name,
                })
        return pages

    @staticmethod
    def _parse_txt(file_path: str, file_name: str) -> List[Dict[str, Any]]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        return [{
            "content": content,
            "page_number": 1,
            "file_name": file_name,
        }]

    @staticmethod
    def detect_file_type(file_name: str) -> str:
        ext = Path(file_name).suffix.lower()
        mapping = {
            ".pdf": "pdf",
            ".docx": "docx",
            ".pptx": "pptx",
            ".txt": "txt",
        }
        return mapping.get(ext, "txt")