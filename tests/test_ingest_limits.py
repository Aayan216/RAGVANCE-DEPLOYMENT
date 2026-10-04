"""Ingestion limits + multi-format + multi-document tests.

Covers: configurable size/count/chunk limits, all supported formats
(PDF, DOCX, PPTX, TXT), multi-document independence, and Practice/Mock
listing of processed documents. Nothing here weakens existing behavior:
successful processing still goes through the same pipeline.
"""
import os
import sys
import time

sys.path.insert(0, r"D:\v5\RAG-p2")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
import django

django.setup()

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client

from backend import views
from backend.models import Chunk, Document

PASSED = 0
FAILED = 0
CREATED = []
ORIG = {}

def ok(label, cond, extra=""):
    global PASSED, FAILED
    if cond:
        PASSED += 1
        print("PASS -", label, flush=True)
    else:
        FAILED += 1
        print("FAIL -", label, (":: " + str(extra)) if extra else "", flush=True)


def set_setting(name, value):
    if name not in ORIG:
        ORIG[name] = getattr(settings, name)
    setattr(settings, name, value)


def upload(title, name, payload, content_type="application/octet-stream"):
    up = SimpleUploadedFile(name, payload, content_type=content_type)
    r = c.post("/", {"title": title, "file": up})
    doc = Document.objects.filter(title=title).first()
    if doc:
        CREATED.append(doc)
    return r, doc


def process_ok(doc, label):
    n0 = len(views.vector_store)
    t0 = time.time()
    r = c.post(f"/process/{doc.id}/")
    sec = time.time() - t0
    data = r.json() if "application/json" in r.headers.get("Content-Type", "") else {}
    ok(f"{label}: process 200 success JSON",
       r.status_code == 200 and data.get("status") == "success"
       and data.get("success") is True and data.get("chunks", 0) > 0,
       (r.status_code, data))
    doc.refresh_from_db()
    ok(f"{label}: chunk_count == rows == FAISS metas",
       doc.chunk_count == data.get("chunks")
       == Chunk.objects.filter(document=doc).count()
       == len([m for m in views.vector_store.metadata
               if m.get("doc_id") == doc.id]) > 0,
       (doc.chunk_count, Chunk.objects.filter(document=doc).count()))
    ok(f"{label}: FAISS grew by chunks",
       len(views.vector_store) == n0 + doc.chunk_count,
       (n0, len(views.vector_store), doc.chunk_count))
    print(f"    [{label}] processed in {sec:.1f}s, chunks={doc.chunk_count}", flush=True)
    return data.get("chunks", 0)


def drop(doc):
    if doc:
        views.vector_store.delete_document(doc.id)
        fp = doc.file.path if doc.file else None
        doc.delete()
        if fp and os.path.exists(fp):
            os.remove(fp)


def diverse_sentences(n, seed=7):
    """n semantically distinct sentences.

    near_duplicate_mask drops chunks with embedding cosine >= 0.97, so
    synthetic fixtures must not reuse uniform phrasing (a whole document of
    near-identical lines collapses to a handful of chunks after dedup).
    """
    from itertools import product
    import random as _random

    nouns = ["estuary", "compressor", "weld seam", "monsoon", "algorithmmodule",
             "ledger", "photosynthesis", "turbine", "porcelain", "canyon",
             "antibody", "ballast", "substrate", "glacier", "cipher",
             "orchard", "reactor", "sediment", "grammar", "chromatography",
             "scaffold", "plankton", "voltmeter", "callus", "keel",
             "stratum", "catalyst", "manuscript", "alluvium", "bearing",
             "lattice", "spore", "catenary", "moraine", "isotope",
             "compass", "veneer", "mycelium", "basalt", "drumlin"]
    verbs = ["reshapes", "measures", "preserves", "constrains", "documents",
             "erodes", "stabilizes", "amplifies", "encodes", "fragments"]
    contexts = ["the shoreline survey", "a brewery fermentation tank",
                "the bridge expansion joint", "winter grazing patterns",
                "the compiler optimizer", "harvest accounting records",
                "chloroplast light capture", "a wind tunnel model",
                "kiln firing tolerances", "the slot canyon hikers",
                "vaccine cold-chain logs", "the ship ballast system",
                "the etched semiconductor", "the polar ice margin",
                "the public key exchange", "the graft union formation",
                "the core sample column", "the reef nutrient cycle",
                "the dialect phoneme inventory", "the amino acid column",
                "the steel reinforcement", "the coral spawning event",
                "the relay contact pitting", "the nursery seedling tray",
                "the propeller blade root", "the ash fall deposit",
                "the batch reaction yield", "the marginalia annotation",
                "the floodplain clay lens", "the gear tooth surface"]
    combos = list(product(nouns, verbs, contexts))
    _random.Random(seed).shuffle(combos)
    return [f"The {n} {v} {x}." for n, v, x in combos[:n]]


c = Client()

try:
    # ---------- A. small TXT ----------
    body_a = ("Alpha signal processing notes for ingest testing.\n\n"
              + " ".join(f"Section {i} covers sampling rates, aliasing filters, "
                        f"and window functions for signal study {i}." for i in range(20)))
    r, doc_a = upload("IngestATxt", "a_small.txt", body_a.encode(), "text/plain")
    ok("A txt upload -> 302", r.status_code == 302 and doc_a is not None, r.status_code)
    if doc_a:
        ok("A file_type txt", doc_a.file_type == "txt", doc_a.file_type)
        process_ok(doc_a, "A txt")
        metas = [m for m in views.vector_store.metadata if m.get("doc_id") == doc_a.id]
        ok("A metadata page_number=1 + filename",
           all(m.get("page_number") == 1 for m in metas)
           and all(m.get("file_name") == "a_small.txt" for m in metas),
           metas[:2])

    # ---------- B. small PDF (2 pages, generated) ----------
    import fitz
    pdf_path = os.path.join(settings.BASE_DIR, "data", "media", "documents",
                             "_gen_small.pdf")
    os.makedirs(os.path.dirname(pdf_path), exist_ok=True)
    gen = fitz.open()
    for p in range(2):
        page = gen.new_page()
        page.insert_text((72, 72),
                         f"Quxbury bench page {p + 1}: calibration tolerance "
                         f"0.03 millimetres across the actuator travel range.")
        page.insert_text((72, 100),
                         f"Logbook entry for page {p + 1} includes operator "
                         f"initials and measured drift values.")
    gen.save(pdf_path)
    gen.close()
    with open(pdf_path, "rb") as fh:
        r, doc_b = upload("IngestBPdf", "b_small.pdf", fh.read(), "application/pdf")
    ok("B pdf upload -> 302", r.status_code == 302 and doc_b is not None, r.status_code)
    if doc_b:
        ok("B file_type pdf", doc_b.file_type == "pdf", doc_b.file_type)
        process_ok(doc_b, "B pdf")
        metas = [m for m in views.vector_store.metadata if m.get("doc_id") == doc_b.id]
        ok("B page numbers preserved (1..2)",
           {m.get("page_number") for m in metas} <= {1, 2}
           and {m.get("page_number") for m in metas} != set(),
           sorted({m.get("page_number") for m in metas}))

    # ---------- C. DOCX (1 logical document) ----------
    from docx import Document as DocxDocument
    docx_path = os.path.join(settings.BASE_DIR, "data", "media", "documents",
                              "_gen_small.docx")
    dx = DocxDocument()
    dx.add_paragraph("Quxbury regulator housing seals are stocked as SL-9 "
                     "with a shelf life of two years.")
    for i in range(15):
        dx.add_paragraph(
            f"Docx paragraph {i} explains coolant pressure checks, error code "
            f"E7 recovery steps, and Class-B sign-off rules for run {i}.")
    dx.save(docx_path)
    with open(docx_path, "rb") as fh:
        r, doc_c = upload("IngestCDocx", "c_small.docx", fh.read(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    ok("C docx upload -> 302", r.status_code == 302 and doc_c is not None, r.status_code)
    if doc_c:
        ok("C file_type docx", doc_c.file_type == "docx", doc_c.file_type)
        process_ok(doc_c, "C docx")
        metas = [m for m in views.vector_store.metadata if m.get("doc_id") == doc_c.id]
        ok("C page_number == 1 (single logical doc)",
           all(m.get("page_number") == 1 for m in metas), metas[:2])

    # ---------- D. PPTX (slide numbers) ----------
    from pptx import Presentation
    pptx_path = os.path.join(settings.BASE_DIR, "data", "media", "documents",
                              "_gen_small.pptx")
    prs = Presentation()
    for s in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Quxbury Slide {s + 1}"
        slide.placeholders[1].text = (
            f"Slide {s + 1} body: calibration interval is every 40 operating "
            f"hours or after power interruptions longer than twelve seconds.")
    prs.save(pptx_path)
    with open(pptx_path, "rb") as fh:
        r, doc_d = upload("IngestDPptx", "d_small.pptx", fh.read(),
                          "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    ok("D pptx upload -> 302", r.status_code == 302 and doc_d is not None, r.status_code)
    if doc_d:
        ok("D file_type pptx", doc_d.file_type == "pptx", doc_d.file_type)
        process_ok(doc_d, "D pptx")
        metas = [m for m in views.vector_store.metadata if m.get("doc_id") == doc_d.id]
        ok("D slide numbers preserved (1..2)",
           {m.get("page_number") for m in metas} <= {1, 2}
           and {m.get("page_number") for m in metas} != set(),
           sorted({m.get("page_number") for m in metas}))

    # ---------- E. large multi-page PDF (80 pages, generated) ----------
    big_path = os.path.join(settings.BASE_DIR, "data", "media", "documents",
                             "_gen_80p.pdf")
    e_lines = diverse_sentences(80 * 10)
    gen = fitz.open()
    for p in range(80):
        page = gen.new_page()
        y = 72
        for line in range(10):
            page.insert_text((72, y),
                             f"Page {p + 1}: " + e_lines[p * 10 + line])
            y += 18
    gen.save(big_path)
    gen.close()
    size_e = os.path.getsize(big_path)
    with open(big_path, "rb") as fh:
        r, doc_e = upload("IngestELargePdf", "e_large_80p.pdf", fh.read(),
                          "application/pdf")
    ok("E 80p upload -> 302", r.status_code == 302 and doc_e is not None,
       (r.status_code, size_e))
    if doc_e:
        chunks_e = process_ok(doc_e, "E 80p pdf")
        metas = [m for m in views.vector_store.metadata if m.get("doc_id") == doc_e.id]
        pages_seen = sorted({m.get("page_number") for m in metas})
        ok("E page numbers span multiple pages <=80",
           len(pages_seen) >= 2 and max(pages_seen) <= 80, pages_seen[:5])
        print(f"    [E] file size {size_e} bytes, chunks={chunks_e}", flush=True)

    # ---------- H. unsupported extension rejected at upload ----------
    count_before = Document.objects.count()
    r, doc_h = upload("IngestHExe", "h_payload.exe", b"MZ\x90\x00fake")
    ok("H unsupported ext -> form rejection (200 re-render)",
       r.status_code == 200 and doc_h is None, r.status_code)
    ok("H visible error message",
       b"Unsupported file type" in r.content, r.content[:0])
    ok("H no doc row created (slot not consumed)",
       Document.objects.count() == count_before,
       Document.objects.count() - count_before)

    # ---------- I/J. size limit (patched to 1 MB for exactness) ----------
    set_setting("MAX_UPLOAD_FILE_SIZE_MB", 1)
    exact = b"x" * (1 * 1024 * 1024)
    r, doc_i = upload("IngestIExact", "i_exact_1mb.txt", exact, "text/plain")
    ok("I exactly-at-limit upload accepted",
       r.status_code == 302 and doc_i is not None, r.status_code)
    over = b"x" * (1 * 1024 * 1024 + 1)
    count_before = Document.objects.count()
    r, doc_j = upload("IngestJOver", "j_over_1mb.txt", over, "text/plain")
    ok("J above-limit upload rejected",
       r.status_code == 200 and doc_j is None, r.status_code)
    ok("J exact message",
       b"File is too large. Maximum allowed size is 1 MB." in r.content,
       "message missing")
    ok("J no doc row (slot not consumed)",
       Document.objects.count() == count_before,
       Document.objects.count() - count_before)

    # ---------- K. document-count limit ----------
    set_setting("MAX_DOCUMENTS", Document.objects.count())
    count_before = Document.objects.count()
    r, doc_k = upload("IngestKLimit", "k_after_limit.txt", b"still here?", "text/plain")
    ok("K upload at limit rejected",
       r.status_code == 200 and doc_k is None, r.status_code)
    ok("K exact message",
       f"Maximum of {count_before} documents allowed.".encode() in r.content,
       "message missing")
    ok("K count unchanged after rejection",
       Document.objects.count() == count_before)
    ok("K existing documents untouched",
       Document.objects.filter(processed=True).count() >= 1)

    # ---------- chunk ceiling (MAX_CHUNKS_PER_DOCUMENT) ----------
    set_setting("MAX_DOCUMENTS", ORIG["MAX_DOCUMENTS"])  # restore before more uploads
    set_setting("MAX_CHUNKS_PER_DOCUMENT", 2)
    # Distinct content per block so near-duplicate dedup keeps them all
    # (identical blocks would be deduped down to <= the ceiling and pass).
    m_lines = diverse_sentences(4 * 12, seed=11)
    blocks = [" ".join(m_lines[b * 12:(b + 1) * 12]) for b in range(4)]
    r, doc_m = upload("IngestMTooMany", "m_too_many_chunks.txt",
                      "\n\n".join(blocks).encode(), "text/plain")
    ok("M ceiling fixture uploaded", doc_m is not None, r.status_code)
    if doc_m:
        n0 = len(views.vector_store)
        r = c.post(f"/process/{doc_m.id}/")
        data = r.json() if "application/json" in r.headers.get("Content-Type", "") else {}
        ok("M over ceiling -> 400 JSON success=false",
           r.status_code == 400 and data.get("success") is False,
           (r.status_code, data))
        ok("M exact guidance message",
           data.get("error") == "Document contains too much content to process safely. "
                                "Please split the document into smaller files.",
           data.get("error"))
        doc_m.refresh_from_db()
        ok("M not marked processed", doc_m.processed is False)
        ok("M no chunk rows", not Chunk.objects.filter(document=doc_m).exists())
        ok("M FAISS untouched (checked before add)",
           len(views.vector_store) == n0, (n0, len(views.vector_store)))

    # ---------- L. multiple documents together ----------
    set_setting("MAX_CHUNKS_PER_DOCUMENT", ORIG["MAX_CHUNKS_PER_DOCUMENT"])
    r, l1 = upload("IngestL1", "l_doc1.txt",
                   ("Lambda cohort notes: the flux regulator spare part is QUX-42Z. "
                    + " ".join(f"L1 detail {i} about calibration cadence and drift "
                              f"logging for entry {i}." for i in range(20))).encode(),
                   "text/plain")
    r, l2 = upload("IngestL2", "l_doc2.txt",
                   ("Mu cohort notes: the alternator spare part is MM-771A. "
                    + " ".join(f"L2 detail {i} about coolant pressure and error "
                              f"E7 handling for entry {i}." for i in range(20))).encode(),
                   "text/plain")
    ok("L two docs uploaded", l1 is not None and l2 is not None)
    if l1 and l2:
        process_ok(l1, "L doc1")
        process_ok(l2, "L doc2")
        ok("L distinct document ids", l1.id != l2.id, (l1.id, l2.id))

        res1 = views.rag_chain._retrieve_context("QUX-42Z spare part", doc_ids=[l1.id])
        res2 = views.rag_chain._retrieve_context("MM-771A spare part", doc_ids=[l2.id])
        ok("L filtered retrieval doc1 works", len(res1) >= 1, len(res1))
        ok("L filtered retrieval doc2 works", len(res2) >= 1, len(res2))
        ok("L no metadata contamination",
           all(x.get("doc_id") == l1.id for x in res1)
           and all(x.get("doc_id") == l2.id for x in res2),
           ([x.get("doc_id") for x in res1], [x.get("doc_id") for x in res2]))

        body = c.get("/practice/").content.decode()
        ok("L practice page lists both processed docs",
           "IngestL1" in body and "IngestL2" in body)
        body = c.get("/mock-test/").content.decode()
        ok("L mock settings page lists both processed docs",
           "IngestL1" in body and "IngestL2" in body)

        # delete doc1 -> doc2 untouched
        n0 = len(views.vector_store)
        count1 = len([m for m in views.vector_store.metadata
                      if m.get("doc_id") == l1.id])
        count2 = len([m for m in views.vector_store.metadata
                      if m.get("doc_id") == l2.id])
        r = c.post(f"/delete/{l1.id}/")
        ok("L delete doc1 -> 200 success",
           r.status_code == 200 and r.json().get("status") == "success",
           (r.status_code, r.content[:150]))
        count2_after = len([m for m in views.vector_store.metadata
                            if m.get("doc_id") == l2.id])
        ok("L doc2 vectors intact after doc1 delete",
           count2_after == count2 and count2 > 0, (count2, count2_after))
        ok("L FAISS shrank by exactly doc1's vectors",
           len(views.vector_store) == n0 - count1 and count1 > 0,
           (n0, count1, len(views.vector_store)))
        ok("L doc2 row still present",
           Document.objects.filter(id=l2.id).exists())
        res2b = views.rag_chain._retrieve_context("MM-771A spare part", doc_ids=[l2.id])
        ok("L doc2 retrieval still works", len(res2b) >= 1, len(res2b))

except Exception as exc:  # noqa: BLE001
    import traceback
    traceback.print_exc()
    ok("test script ran without unexpected crash", False, repr(exc))

finally:
    for name, value in ORIG.items():
        setattr(settings, name, value)
    for d in CREATED:
        try:
            fresh = Document.objects.filter(id=d.id).first()
            if fresh:
                drop(fresh)
        except Exception as exc:  # noqa: BLE001
            print(f"[cleanup] {d.id}: {exc}", flush=True)
    for gen_file in ("_gen_small.pdf", "_gen_small.docx", "_gen_small.pptx",
                     "_gen_80p.pdf"):
        p = os.path.join(settings.BASE_DIR, "data", "media", "documents", gen_file)
        if os.path.exists(p):
            os.remove(p)
    print("[cleanup] ingest fixtures removed", flush=True)

print(f"\n{PASSED}/{PASSED + FAILED} passed", flush=True)
sys.exit(1 if FAILED else 0)
